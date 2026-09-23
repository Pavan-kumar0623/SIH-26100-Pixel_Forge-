"""
Export Service
Generates CSV (pandas) and Excel (openpyxl) exports.
"""
import io
from sqlalchemy.orm import Session
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models.bidder import Bidder
from models.requirement import Requirement
from models.compliance_result import ComplianceResult
from models.risk_signal import RiskSignal
from models.document import Document
from models.extracted_field import ExtractedField

STATUS_COLORS = {
    "VERIFIED": "C6EFCE",
    "MISSING": "FFC7CE",
    "REVIEW": "FFEB9C",
    "NOT_AVAILABLE": "D9D9D9",
    "HIGH": "FFC7CE",
    "MEDIUM": "FFEB9C",
    "LOW": "C6EFCE",
    "OPEN": "FFC7CE",
    "RESOLVED": "C6EFCE",
}


def export_csv(tender_id: int, db: Session) -> bytes:
    """Generate CSV export using pandas."""
    bidders = db.query(Bidder).filter(Bidder.tender_id == tender_id).all()
    requirements = db.query(Requirement).filter(Requirement.tender_id == tender_id).all()

    rows = []
    for bidder in bidders:
        for req in requirements:
            compliance = db.query(ComplianceResult).filter(
                ComplianceResult.bidder_id == bidder.id,
                ComplianceResult.requirement_id == req.id
            ).first()

            risks = db.query(RiskSignal).filter(
                RiskSignal.bidder_id == bidder.id
            ).all()

            rows.append({
                "Bidder": bidder.company_name,
                "GSTIN": bidder.gstin or "",
                "PAN": bidder.pan or "",
                "Requirement": req.name,
                "Mandatory": "Yes" if req.mandatory else "No",
                "Expected Document": req.expected_document_type or "",
                "Compliance Status": compliance.status if compliance else "MISSING",
                "Reason": compliance.reason if compliance else "",
                "Evidence": compliance.evidence if compliance else "",
                "High Risks": sum(1 for r in risks if r.severity == "HIGH" and r.status == "OPEN"),
                "Medium Risks": sum(1 for r in risks if r.severity == "MEDIUM" and r.status == "OPEN"),
                "Low Risks": sum(1 for r in risks if r.severity == "LOW" and r.status == "OPEN"),
                "Relationship Signals": sum(1 for r in risks if r.risk_type in ["SHARED_IDENTIFIER", "SHARED_ADDRESS"]),
            })

    df = pd.DataFrame(rows)
    buffer = io.BytesIO()
    df.to_csv(buffer, index=False)
    return buffer.getvalue()


def export_excel(tender_id: int, db: Session) -> bytes:
    """Generate multi-sheet Excel export using openpyxl."""
    bidders = db.query(Bidder).filter(Bidder.tender_id == tender_id).all()
    requirements = db.query(Requirement).filter(Requirement.tender_id == tender_id).all()

    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    _write_comparison_sheet(wb, bidders, requirements, db)
    _write_document_sheet(wb, bidders, db)
    _write_compliance_sheet(wb, bidders, requirements, db)
    _write_risk_sheet(wb, bidders, db)
    _write_relationship_sheet(wb, bidders, db)
    _write_evidence_sheet(wb, bidders, db)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _header_row(ws, headers: list, fill_color: str = "1F3864"):
    """Write a styled header row."""
    row = [ws.cell(row=1, column=i + 1, value=h) for i, h in enumerate(headers)]
    for cell in row:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22


def _color_cell(cell, status: str):
    color = STATUS_COLORS.get(status)
    if color:
        cell.fill = PatternFill("solid", fgColor=color)


def _write_comparison_sheet(wb, bidders, requirements, db):
    ws = wb.create_sheet("Bidder Comparison")
    headers = ["Requirement", "Mandatory"] + [b.company_name for b in bidders]
    _header_row(ws, headers)

    for row_idx, req in enumerate(requirements, start=2):
        ws.cell(row=row_idx, column=1, value=req.name)
        ws.cell(row=row_idx, column=2, value="Yes" if req.mandatory else "No")
        for col_idx, bidder in enumerate(bidders, start=3):
            result = db.query(ComplianceResult).filter(
                ComplianceResult.bidder_id == bidder.id,
                ComplianceResult.requirement_id == req.id
            ).first()
            status = result.status if result else "MISSING"
            cell = ws.cell(row=row_idx, column=col_idx, value=status)
            _color_cell(cell, status)

    _auto_size(ws)


def _write_document_sheet(wb, bidders, db):
    ws = wb.create_sheet("Document Verification")
    headers = ["Bidder", "File Name", "Document Type", "Status", "OCR Used", "Confidence"]
    _header_row(ws, headers)
    row_idx = 2
    for bidder in bidders:
        docs = db.query(Document).filter(Document.bidder_id == bidder.id).all()
        for doc in docs:
            ws.cell(row=row_idx, column=1, value=bidder.company_name)
            ws.cell(row=row_idx, column=2, value=doc.file_name)
            ws.cell(row=row_idx, column=3, value=doc.document_type)
            status_cell = ws.cell(row=row_idx, column=4, value=doc.status)
            _color_cell(status_cell, doc.status)
            ws.cell(row=row_idx, column=5, value=doc.ocr_used)
            ws.cell(row=row_idx, column=6, value=f"{(doc.classification_confidence or 0):.0%}")
            row_idx += 1
    _auto_size(ws)


def _write_compliance_sheet(wb, bidders, requirements, db):
    ws = wb.create_sheet("Compliance Results")
    headers = ["Bidder", "Requirement", "Mandatory", "Status", "Reason", "Evidence"]
    _header_row(ws, headers)
    row_idx = 2
    for bidder in bidders:
        for req in requirements:
            result = db.query(ComplianceResult).filter(
                ComplianceResult.bidder_id == bidder.id,
                ComplianceResult.requirement_id == req.id
            ).first()
            ws.cell(row=row_idx, column=1, value=bidder.company_name)
            ws.cell(row=row_idx, column=2, value=req.name)
            ws.cell(row=row_idx, column=3, value="Yes" if req.mandatory else "No")
            status = result.status if result else "MISSING"
            status_cell = ws.cell(row=row_idx, column=4, value=status)
            _color_cell(status_cell, status)
            ws.cell(row=row_idx, column=5, value=result.reason if result else "")
            ws.cell(row=row_idx, column=6, value=result.evidence if result else "")
            row_idx += 1
    _auto_size(ws)


def _write_risk_sheet(wb, bidders, db):
    ws = wb.create_sheet("Risk Signals")
    headers = ["Bidder", "Risk Type", "Severity", "Description", "Source", "Evidence", "Status"]
    _header_row(ws, headers)
    row_idx = 2
    for bidder in bidders:
        risks = db.query(RiskSignal).filter(RiskSignal.bidder_id == bidder.id).all()
        for risk in risks:
            ws.cell(row=row_idx, column=1, value=bidder.company_name)
            ws.cell(row=row_idx, column=2, value=risk.risk_type)
            sev_cell = ws.cell(row=row_idx, column=3, value=risk.severity)
            _color_cell(sev_cell, risk.severity)
            ws.cell(row=row_idx, column=4, value=risk.description)
            ws.cell(row=row_idx, column=5, value=risk.source)
            ws.cell(row=row_idx, column=6, value=risk.evidence)
            status_cell = ws.cell(row=row_idx, column=7, value=risk.status)
            _color_cell(status_cell, risk.status)
            row_idx += 1
    _auto_size(ws)


def _write_relationship_sheet(wb, bidders, db):
    ws = wb.create_sheet("Relationship Analysis")
    headers = ["Bidder", "Relationship Type", "Description", "Evidence", "Severity"]
    _header_row(ws, headers)
    row_idx = 2
    for bidder in bidders:
        signals = db.query(RiskSignal).filter(
            RiskSignal.bidder_id == bidder.id,
            RiskSignal.risk_type.in_(["SHARED_IDENTIFIER", "SHARED_ADDRESS", "SHARED_DIRECTOR", "RELATED_COMPANY"])
        ).all()
        for s in signals:
            ws.cell(row=row_idx, column=1, value=bidder.company_name)
            ws.cell(row=row_idx, column=2, value=s.risk_type)
            ws.cell(row=row_idx, column=3, value=s.description)
            ws.cell(row=row_idx, column=4, value=s.evidence)
            sev_cell = ws.cell(row=row_idx, column=5, value=s.severity)
            _color_cell(sev_cell, s.severity)
            row_idx += 1
    _auto_size(ws)


def _write_evidence_sheet(wb, bidders, db):
    ws = wb.create_sheet("Evidence")
    headers = ["Bidder", "Document", "Field Name", "Field Value", "Confidence", "Validation", "Evidence"]
    _header_row(ws, headers)
    row_idx = 2
    for bidder in bidders:
        docs = db.query(Document).filter(Document.bidder_id == bidder.id).all()
        for doc in docs:
            fields = db.query(ExtractedField).filter(ExtractedField.document_id == doc.id).all()
            for f in fields:
                ws.cell(row=row_idx, column=1, value=bidder.company_name)
                ws.cell(row=row_idx, column=2, value=doc.file_name)
                ws.cell(row=row_idx, column=3, value=f.field_name)
                ws.cell(row=row_idx, column=4, value=f.field_value)
                ws.cell(row=row_idx, column=5, value=f"{f.confidence:.0%}")
                val_cell = ws.cell(row=row_idx, column=6, value=f.validation_status)
                _color_cell(val_cell, f.validation_status)
                ws.cell(row=row_idx, column=7, value=f.evidence)
                row_idx += 1
    _auto_size(ws)


def _auto_size(ws):
    """Auto-size columns."""
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 4, 60)
